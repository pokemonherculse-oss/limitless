"""Smoke test GitHub Excel persistence, restart recovery and ephemeral chat.

Runs with the standard library (plus the project's openpyxl dependency).
It mocks the GitHub Contents API; it does NOT need GitHub credentials.
"""
import base64
import hashlib
import importlib
import json
import os
import tempfile
import threading
import urllib.error
import urllib.request
from http.cookiejar import CookieJar
from http.server import ThreadingHTTPServer
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(BASE))

with tempfile.TemporaryDirectory() as work:
    os.environ['GITHUB_DATA_REPO'] = 'tester/limitless-data'
    os.environ['GITHUB_TOKEN'] = 'fake-token'
    os.environ['ADMIN_USER'] = 'admin'
    os.environ['ADMIN_PASSWORD'] = 'check-only-password'
    os.environ['SESSION_SECRET'] = 'test-fixed-secret'
    os.environ['LIMITLESS_EXCEL_PATH'] = str(Path(work) / 'local.xlsx')

    import github_excel
    data_path = Path(work) / 'github.xlsx'
    calls = []

    def fake_api(self, path, method='GET', payload=None):
        calls.append((path, method))
        if path == '/repos/tester/limitless-data':
            return 200, {'full_name': 'tester/limitless-data'}
        if path.startswith('/repos/tester/limitless-data/contents/') and method == 'GET':
            if not data_path.exists():
                return 404, None
            raw = data_path.read_bytes()
            return 200, {'type':'file', 'sha':hashlib.sha1(raw).hexdigest(), 'encoding':'base64', 'content':base64.b64encode(raw).decode()}
        if path == '/repos/tester/limitless-data/contents/limitless.xlsx' and method == 'PUT':
            raw = base64.b64decode(payload['content'])
            if data_path.exists():
                assert payload['sha'] == hashlib.sha1(data_path.read_bytes()).hexdigest(), 'Wrong GitHub revision'
            else:
                assert 'sha' not in payload
            data_path.write_bytes(raw)
            return 200, {'content': {'sha': hashlib.sha1(raw).hexdigest()}}
        raise AssertionError(f'Unexpected API call {path} {method}')

    github_excel.GitHubExcel._request = fake_api
    import server
    assert data_path.exists(), 'initial workbook must be initialized remotely'

    httpd = ThreadingHTTPServer(('127.0.0.1',0), server.Handler)
    base = 'http://127.0.0.1:' + str(httpd.server_port)
    worker = threading.Thread(target=httpd.serve_forever, daemon=True)
    worker.start()

    def client():
        return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(CookieJar()))

    def request(opener, path, body=None):
        raw = None if body is None else json.dumps(body).encode()
        rq = urllib.request.Request(base+path, data=raw, headers={'Content-Type':'application/json'} if raw else {})
        with opener.open(rq,timeout=6) as f:
            if path == '/api/export':
                return f.read()
            return json.load(f)

    alice = client()
    bob = client()
    admin = client()
    assert request(alice,'/api/register',{'username':'Alice','password':'abc1234'})['ok']
    assert request(bob,'/api/register',{'username':'Bob123','password':'abc1234'})['ok']
    assert request(admin,'/api/admin/login',{'username':'admin','password':'check-only-password'})['ok']
    state = request(admin,'/api/state')['state']
    assert len(state['players']) == 2
    alice_id, bob_id = [p['id'] for p in state['players']]
    tour = {'id':'test-tournament','name':'Tomorrow battle','status':'ACTIVE',
            'pairings':[{'round':1,'processed':False,
                        'matches':[{'matchId':'battle1','p1':alice_id,'p2':bob_id}]}],
            'registeredPlayers':[alice_id,bob_id]}
    out = request(admin,'/api/patch', {'ops':[{'op':'set','path':['tournaments',{'id':'test-tournament'}],'value':tour}]})
    assert out['revision'] >= 3
    assert request(alice,'/api/chat/send',{'matchId':'battle1','text':'Battle ID 123-XYZ'})['ok']
    state_bob = request(bob,'/api/state')['state']
    msgs = state_bob['tournaments'][0]['pairings'][0]['matches'][0].get('chat',[])
    assert any('123-XYZ' in m['text'] for m in msgs)
    from io import BytesIO
    from openpyxl import load_workbook
    export = request(admin,'/api/export')
    wb = load_workbook(BytesIO(export),read_only=True)
    assert 'Records' in wb.sheetnames
    assert not any('123-XYZ' in str(row) for row in wb['Records'].values), 'Chat accidentally persisted in Excel'
    wb.close()

    # Fake Render restart: remove the *entire* local Excel cache, reload Python
    # module and require restoration of players/tournament from private GitHub.
    httpd.shutdown()
    httpd.server_close()
    Path(os.environ['LIMITLESS_EXCEL_PATH']).unlink()
    server = importlib.reload(server)
    assert len(server.STATE['players']) == 2
    assert len(server.STATE['tournaments']) == 1
    assert server.CHAT == {}, 'Chat must be gone after restart'
    assert Path(os.environ['LIMITLESS_EXCEL_PATH']).exists()
    assert request is not None
    print('PASS: player registration, admin tournament creation, cross-player chat, no chat in Excel, GitHub recovery after Render restart')
